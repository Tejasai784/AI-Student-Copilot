import io
from pathlib import Path
from typing import List, Dict, Union, Optional
import pypdf
import docx
from backend.logging_config import logger


class DocumentExtractionError(Exception):
    """Custom exception raised when document parsing fails."""
    pass


# Backwards compatibility alias
PDFExtractionError = DocumentExtractionError


def extract_text_from_pdf(source: Union[str, Path, bytes, io.BytesIO]) -> List[Dict[str, Union[int, str]]]:
    """
    Extracts text page-by-page from a PDF file while preserving page numbers.

    Args:
        source: File path, raw bytes, or BytesIO stream of the PDF.

    Returns:
        List of dicts: [{"page_number": 1, "text": "..."}, ...]

    Raises:
        DocumentExtractionError: If file is corrupted, encrypted, or invalid.
    """
    stream = None
    try:
        if isinstance(source, (str, Path)):
            stream = open(source, "rb")
        elif isinstance(source, bytes):
            stream = io.BytesIO(source)
        else:
            stream = source

        reader = pypdf.PdfReader(stream)
        
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise DocumentExtractionError("PDF file is password protected and cannot be processed.")

        total_pages = len(reader.pages)
        if total_pages == 0:
            raise DocumentExtractionError("PDF contains 0 pages.")

        pages_data = []
        total_chars = 0

        for idx, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            pages_data.append({
                "page_number": idx,
                "text": page_text
            })
            total_chars += len(page_text.strip())

        logger.info(f"Extracted {total_pages} pages ({total_chars} characters) from PDF.")
        return pages_data

    except DocumentExtractionError:
        raise
    except Exception as e:
        logger.error(f"Failed to extract text from PDF: {e}", exc_info=True)
        raise DocumentExtractionError(f"Corrupt or invalid PDF file: {str(e)}")
    finally:
        if isinstance(source, (str, Path)) and stream and not stream.closed:
            stream.close()


def extract_text_from_docx(source: Union[str, Path, bytes, io.BytesIO]) -> List[Dict[str, Union[int, str]]]:
    """Extracts text from a DOCX document, grouped into virtual pages."""
    stream = None
    try:
        if isinstance(source, (str, Path)):
            stream = open(source, "rb")
        elif isinstance(source, bytes):
            stream = io.BytesIO(source)
        else:
            stream = source

        doc = docx.Document(stream)
        paragraphs_text = []
        for p in doc.paragraphs:
            txt = p.text.strip()
            if txt:
                paragraphs_text.append(txt)

        # Include tables
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    paragraphs_text.append(" | ".join(row_cells))

        full_text = "\n\n".join(paragraphs_text)
        if not full_text.strip():
            raise DocumentExtractionError("DOCX document contains no readable text.")

        # Chunk into virtual pages of roughly 2500 characters
        pages = []
        page_size = 2500
        total_len = len(full_text)
        for idx, start in enumerate(range(0, total_len, page_size), start=1):
            pages.append({
                "page_number": idx,
                "text": full_text[start:start + page_size]
            })

        logger.info(f"Extracted {len(pages)} virtual pages from DOCX ({len(full_text)} characters).")
        return pages
    except DocumentExtractionError:
        raise
    except Exception as e:
        logger.error(f"Failed to extract text from DOCX: {e}", exc_info=True)
        raise DocumentExtractionError(f"Corrupt or invalid DOCX file: {str(e)}")
    finally:
        if isinstance(source, (str, Path)) and stream and not stream.closed:
            stream.close()


def extract_text_from_txt(source: Union[str, Path, bytes, io.BytesIO]) -> List[Dict[str, Union[int, str]]]:
    """Extracts text from TXT or MD file, grouped into virtual pages."""
    try:
        if isinstance(source, (str, Path)):
            with open(source, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
        elif isinstance(source, bytes):
            content = source.decode("utf-8", errors="replace")
        elif isinstance(source, io.BytesIO):
            content = source.getvalue().decode("utf-8", errors="replace")
        else:
            content = source.read()
            if isinstance(content, bytes):
                content = content.decode("utf-8", errors="replace")

        if not content.strip():
            raise DocumentExtractionError("Text document is empty.")

        pages = []
        page_size = 2500
        total_len = len(content)
        for idx, start in enumerate(range(0, total_len, page_size), start=1):
            pages.append({
                "page_number": idx,
                "text": content[start:start + page_size]
            })

        logger.info(f"Extracted {len(pages)} virtual pages from text file ({total_len} characters).")
        return pages
    except DocumentExtractionError:
        raise
    except Exception as e:
        logger.error(f"Failed to extract text from TXT: {e}", exc_info=True)
        raise DocumentExtractionError(f"Invalid text file: {str(e)}")


def extract_text_from_csv(source: Union[str, Path, bytes, io.BytesIO]) -> List[Dict[str, Union[int, str]]]:
    """Extracts text representation from CSV file."""
    import pandas as pd
    try:
        if isinstance(source, (str, Path)):
            df = pd.read_csv(source)
        elif isinstance(source, bytes):
            df = pd.read_csv(io.BytesIO(source))
        else:
            df = pd.read_csv(source)

        if df.empty:
            raise DocumentExtractionError("CSV file contains no data rows.")

        summary_lines = [
            f"CSV Summary: {len(df)} rows, {len(df.columns)} columns: {', '.join(df.columns.astype(str))}",
            "\nSample Data:\n" + df.head(50).to_string(index=False)
        ]
        text_content = "\n".join(summary_lines)

        return [{"page_number": 1, "text": text_content}]
    except DocumentExtractionError:
        raise
    except Exception as e:
        logger.error(f"Failed to extract CSV: {e}", exc_info=True)
        raise DocumentExtractionError(f"Invalid CSV file: {str(e)}")


def extract_text_from_file(source: Union[str, Path, bytes, io.BytesIO], filename: str) -> List[Dict[str, Union[int, str]]]:
    """
    Universal extractor that dispatches based on filename extension:
    .pdf -> extract_text_from_pdf
    .docx -> extract_text_from_docx
    .txt, .md -> extract_text_from_txt
    .csv -> extract_text_from_csv
    """
    fn = filename.lower()
    if fn.endswith(".pdf"):
        return extract_text_from_pdf(source)
    elif fn.endswith(".docx"):
        return extract_text_from_docx(source)
    elif fn.endswith(".txt") or fn.endswith(".md") or fn.endswith(".json"):
        return extract_text_from_txt(source)
    elif fn.endswith(".csv"):
        return extract_text_from_csv(source)
    else:
        # Default try PDF or text
        try:
            return extract_text_from_pdf(source)
        except Exception:
            return extract_text_from_txt(source)
